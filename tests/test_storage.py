from specrec import storage


def test_csv_and_xlsx(tmp_path, monkeypatch):
    monkeypatch.setenv("SPECREC_DATA", str(tmp_path))
    p = storage.day_dir() / "photometry.csv"
    storage.append_row(p, ["a", "b"], [1, 0.5])
    storage.append_row(p, ["a", "b"], [2, float("inf")])
    lines = p.read_text(encoding="utf-8-sig").splitlines()
    assert lines == ["a,b", "1,0.5000", "2,inf"]
    x = storage.write_table(tmp_path / "t.xlsx", ["λ", "A"], [[500, 0.1], [510, None]], {"sample": "s"})
    from openpyxl import load_workbook
    ws = load_workbook(x).active
    assert ws["A3"].value == "λ" and ws["B4"].value == 0.1


def test_append_row_new_file_when_header_changes(tmp_path):
    p = tmp_path / "photometry.csv"
    storage.append_row(p, ["a", "b"], [1, 2])
    storage.append_row(p, ["a", "b", "cuvette"], [3, 4, "скло"])
    storage.append_row(p, ["a", "b", "cuvette"], [5, 6, "кварц"])
    assert p.read_text(encoding="utf-8-sig").splitlines() == ["a,b", "1,2"]
    p2 = tmp_path / "photometry_2.csv"
    assert p2.read_text(encoding="utf-8-sig").splitlines() == ["a,b,cuvette", "3,4,скло", "5,6,кварц"]
