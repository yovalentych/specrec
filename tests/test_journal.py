import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from specrec.journal import write_runs
from specrec.protocol import Run
from specrec.protocols_builtin import SOD, SH

from test_protocol import finish


def make_journal(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Зразки"
    ws["B4"], ws["C4"] = "Код тварини", "Група"
    ws["B5"], ws["C5"] = "C1", "К"
    ws["B6"], ws["C6"] = "D1", "DOX"
    sod = wb.create_sheet("СОД")
    sod["I17"] = '=IF(OR(E17="",F17=""),"",F17-E17)'      # a formula must survive
    wb.create_sheet("SH")
    wb.save(path)


def test_write_sod_and_sh(tmp_path):
    j = tmp_path / "journal.xlsx"
    make_journal(j)
    samples = [{"code": "D1", "group": "DOX"}, {"code": "C1", "group": "К"}, {"code": "NEW", "group": "К"}]
    sod = Run(SOD, samples)
    finish(sod, lambda t, r, wl: {"control": {"A0": 0.1, "A10": 0.6}, "test": {"A0": 0.11, "A10": 0.35}}
           .get(t.type, {}).get(r, 0.0))
    sh = Run(SH, samples[:2])
    finish(sh, lambda t, r, wl: 0.0 if t.type == "zero" else {380: 0.05, 412: 0.07, 415: 0.08, 450: 0.04}.get(wl, 0.3)
           if r == "fon" else 0.3)
    out = tmp_path / "out.xlsx"
    report = write_runs(j, out, [sod, sh])
    wb = load_workbook(out)
    s = wb["СОД"]
    assert s["B8"].value == 0.1 and s["C10"].value == 0.6
    assert s["E17"].value == 0.11 and s["H17"].value == 0.35       # C1 is row 5 -> 17
    assert s["E18"].value == 0.11                                  # D1 is row 6 -> 18
    assert s["I17"].value.startswith("=IF")                        # formula kept
    assert wb["Зразки"]["B7"].value == "NEW" and s["E19"].value == 0.11
    sh_ws = wb["SH"]
    assert sh_ws["D12"].value == 0.07 and wb["Зразки"]["K5"].value == 0.08
    assert any("СОД" in line for line in report)
    assert load_workbook(j)["СОД"]["E17"].value is None            # the original is untouched
