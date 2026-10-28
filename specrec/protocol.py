"""Protocols: a guided series of tubes, timers and formulas (no GUI here).

A protocol (a plain dict, stored as JSON for custom ones) describes:
  tubes        tube types: label, how many per series or per sample, the steps of one tube
               steps: {"do": text}  an action the operator confirms
                      {"wait": seconds, "from": i, "text": ...}  a timer (from step i, default: previous)
                      {"read": name, "wl": [nm, ...]}  a measurement at one or more wavelengths
               the tube type with "zero": true sets the zero of the instrument (all its wavelengths)
  constants    numbers used in formulas (volumes, ε, times)
  inputs       per-sample numbers the operator can change (dilution …) with defaults
  calibration  optional: standards of one tube type with known x → slope/intercept/R²
  tube_calc    per tube type: ordered {var: formula}
  series_calc  ordered {var: formula} over series tubes (e.g. mean of controls)
  sample_calc  ordered {var: formula} over the tubes of one sample
               (`test` = the measured values of the sample's "test" tubes, `test.dA` = a computed one)
  results      what goes into the results table
  qc           checks: {scope: series|sample, expr, min, max, msg}

Formulas are evaluated safely (a small subset of Python: numbers, + − * / **, comparisons,
`a if c else b`, functions below). In a sample formula a tube type name is the list of its
values (`mean(test)`), `test.dA` the list of a computed variable; a tube's reads are
`A0`, `A10` (one wavelength) or `fon_412` (`<read>_<wl>`).
"""
from __future__ import annotations

import ast
import copy
import datetime as dt
import json
import math
import time
from dataclasses import dataclass, field
from pathlib import Path

# ---------------------------------------------------------------- safe formulas
NAN = float("nan")


def _clean(xs):
    if isinstance(xs, (int, float)):
        xs = [xs]
    return [float(x) for x in xs if x is not None and not (isinstance(x, float) and math.isnan(x))]


def _mean(xs):
    v = _clean(xs)
    return sum(v) / len(v) if v else NAN


def _sd(xs):
    v = _clean(xs)
    if len(v) < 2:
        return NAN
    m = sum(v) / len(v)
    return math.sqrt(sum((x - m) ** 2 for x in v) / (len(v) - 1))


def _cv(xs):
    m = _mean(xs)
    return _sd(xs) / abs(m) if m and not math.isnan(m) else NAN


FUNCS = {
    "mean": _mean, "sd": _sd, "cv": _cv,
    "n": lambda xs: float(len(_clean(xs))),
    "sum": lambda xs: float(sum(_clean(xs))),
    "min": lambda *a: min(_clean(a[0] if len(a) == 1 and isinstance(a[0], list) else a), default=NAN),
    "max": lambda *a: max(_clean(a[0] if len(a) == 1 and isinstance(a[0], list) else a), default=NAN),
    "first": lambda xs: (_clean(xs) or [NAN])[0],
    "abs": abs, "sqrt": math.sqrt, "log10": math.log10, "ln": math.log, "exp": math.exp,
    "has": lambda x: not (isinstance(x, float) and math.isnan(x)),
    "nz": lambda x, d=0.0: d if isinstance(x, float) and math.isnan(x) else x,   # NaN -> default
}

_OPS = {ast.Add: lambda a, b: a + b, ast.Sub: lambda a, b: a - b, ast.Mult: lambda a, b: a * b,
        ast.Div: lambda a, b: a / b if b != 0 else NAN, ast.Pow: lambda a, b: a ** b}
_CMP = {ast.Lt: lambda a, b: a < b, ast.LtE: lambda a, b: a <= b, ast.Gt: lambda a, b: a > b,
        ast.GtE: lambda a, b: a >= b, ast.Eq: lambda a, b: a == b, ast.NotEq: lambda a, b: a != b}


class FormulaError(Exception):
    pass


def evaluate(expr: str, names: dict):
    """Evaluate a formula with the given names. Unknown names are NaN (= not measured yet)."""
    try:
        tree = ast.parse(str(expr), mode="eval")
    except SyntaxError as e:
        raise FormulaError(f"Помилка у формулі «{expr}»: {e.msg}") from None

    def ev(node):
        if isinstance(node, ast.Expression):
            return ev(node.body)
        if isinstance(node, ast.Constant) and isinstance(node.value, (int, float)):
            return float(node.value)
        if isinstance(node, ast.Name):
            return names.get(node.id, NAN)
        if isinstance(node, ast.Attribute) and isinstance(node.value, ast.Name):
            return names.get(f"{node.value.id}.{node.attr}", [])
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.USub, ast.UAdd)):
            v = ev(node.operand)
            return -v if isinstance(node.op, ast.USub) else v
        if isinstance(node, ast.BinOp) and type(node.op) in _OPS:
            a, b = ev(node.left), ev(node.right)
            if isinstance(a, list) or isinstance(b, list):
                raise FormulaError(f"У «{expr}» список (кілька пробірок) треба обгорнути в mean(...)")
            try:
                return _OPS[type(node.op)](a, b)
            except (OverflowError, ValueError, ZeroDivisionError):
                return NAN
        if isinstance(node, ast.Compare) and len(node.ops) == 1 and type(node.ops[0]) in _CMP:
            return _CMP[type(node.ops[0])](ev(node.left), ev(node.comparators[0]))
        if isinstance(node, ast.BoolOp):
            vals = [ev(v) for v in node.values]
            return all(vals) if isinstance(node.op, ast.And) else any(vals)
        if isinstance(node, ast.IfExp):
            return ev(node.body) if ev(node.test) else ev(node.orelse)
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and (
                node.func.id in FUNCS or callable(names.get(node.func.id))):
            args = [ev(a) for a in node.args]
            fn = names[node.func.id] if callable(names.get(node.func.id)) else FUNCS[node.func.id]
            try:
                return fn(*args)
            except (ValueError, TypeError, ZeroDivisionError, OverflowError):
                return NAN
        raise FormulaError(f"У формулі «{expr}» не дозволено: {ast.dump(node)[:60]}")

    return ev(tree)


def isnum(x) -> bool:
    return isinstance(x, (int, float)) and not isinstance(x, bool) and not math.isnan(x)


# ---------------------------------------------------------------- statistics
def _betacf(a, b, x):
    qab, qap, qam = a + b, a + 1, a - 1
    c, d = 1.0, 1 - qab * x / qap
    d = 1 / (d if abs(d) > 1e-30 else 1e-30)
    h = d
    for m in range(1, 300):
        m2 = 2 * m
        aa = m * (b - m) * x / ((qam + m2) * (a + m2))
        d = 1 + aa * d
        d = 1 / (d if abs(d) > 1e-30 else 1e-30)
        c = 1 + aa / c if abs(c) > 1e-30 else 1e30
        h *= d * c
        aa = -(a + m) * (qab + m) * x / ((a + m2) * (qap + m2))
        d = 1 + aa * d
        d = 1 / (d if abs(d) > 1e-30 else 1e-30)
        c = 1 + aa / c if abs(c) > 1e-30 else 1e30
        de = d * c
        h *= de
        if abs(de - 1) < 1e-12:
            break
    return h


def betainc(a, b, x):
    """Regularized incomplete beta I_x(a, b)."""
    if x <= 0:
        return 0.0
    if x >= 1:
        return 1.0
    lbt = math.lgamma(a + b) - math.lgamma(a) - math.lgamma(b) + a * math.log(x) + b * math.log(1 - x)
    if x < (a + 1) / (a + b + 2):
        return math.exp(lbt) * _betacf(a, b, x) / a
    return 1 - math.exp(lbt) * _betacf(b, a, 1 - x) / b


def t_two_sided_p(t, df):
    return betainc(df / 2, 0.5, df / (df + t * t)) if df > 0 else NAN


def welch(x, y):
    """Welch t-test: (t, df, p two-sided) for y vs x."""
    x, y = _clean(x), _clean(y)
    if len(x) < 2 or len(y) < 2:
        return NAN, NAN, NAN
    vx, vy = _sd(x) ** 2 / len(x), _sd(y) ** 2 / len(y)
    if vx + vy == 0:
        return NAN, NAN, NAN
    t = (_mean(y) - _mean(x)) / math.sqrt(vx + vy)
    df = (vx + vy) ** 2 / (vx ** 2 / (len(x) - 1) + vy ** 2 / (len(y) - 1))
    return t, df, t_two_sided_p(abs(t), df)


# ---------------------------------------------------------------- run state
@dataclass
class Tube:
    uid: str
    type: str
    label: str
    sample: str | None = None        # sample code, None for series tubes
    x: float | None = None           # known value of a standard
    step: int = 0                    # index of the current step
    reads: dict = field(default_factory=dict)      # "<read>_<wl>" -> A
    done_at: dict = field(default_factory=dict)    # step index -> epoch seconds
    due: float | None = None         # epoch seconds when the current wait step ends

    def to_dict(self):
        return self.__dict__.copy()


class Run:
    """One series of a protocol: series tubes + tubes of every sample, with their progress."""

    def __init__(self, protocol: dict, samples: list[dict]):
        self.p = copy.deepcopy(protocol)
        self.started = dt.datetime.now()
        self.samples = [dict(code=s["code"], group=s.get("group", ""),
                             inputs={k: float(s.get("inputs", {}).get(k, v.get("default", 1)))
                                     for k, v in self.p.get("inputs", {}).items()})
                        for s in samples]
        self.tubes: list[Tube] = []
        self.refs_ready = False       # the zero was measured in this program session
        self.notes = ""
        self._build()

    # ----- structure
    def tube_def(self, t: str) -> dict:
        return self.p["tubes"][t]

    def steps(self, tube: Tube) -> list[dict]:
        return self.tube_def(tube.type)["steps"]

    def _build(self):
        n = 0
        cal = self.p.get("calibration")
        for t, d in self.p["tubes"].items():
            if d.get("per") == "sample":
                continue
            if cal and t == cal["type"]:
                for i, x in enumerate(cal["x"], 1):
                    n += 1
                    self.tubes.append(Tube(f"s{n}", t, d["label"].format(i=i, x=_fmt_x(x)), x=float(x)))
            else:
                for i in range(1, int(d.get("count", 1)) + 1):
                    n += 1
                    lab = d["label"].format(i=i) if "{i}" in d["label"] else (
                        d["label"] + (f" {i}" if int(d.get("count", 1)) > 1 else ""))
                    self.tubes.append(Tube(f"s{n}", t, lab))
        for s in self.samples:
            for t, d in self.p["tubes"].items():
                if d.get("per") != "sample":
                    continue
                cnt = int(d.get("count", 1))
                for i in range(1, cnt + 1):
                    n += 1
                    lab = d["label"].format(i=i) if "{i}" in d["label"] else (d["label"] + (f" {i}" if cnt > 1 else ""))
                    self.tubes.append(Tube(f"t{n}", t, lab, sample=s["code"]))

    def zero_tube(self) -> Tube | None:
        return next((t for t in self.tubes if self.tube_def(t.type).get("zero")), None)

    def wavelengths(self) -> list[int]:
        wls = set()
        for d in self.p["tubes"].values():
            for st in d["steps"]:
                wls.update(st.get("wl", []))
        return sorted(wls)

    # ----- progress
    def current(self, tube: Tube) -> dict | None:
        st = self.steps(tube)
        return st[tube.step] if tube.step < len(st) else None

    def state(self, tube: Tube, now: float | None = None) -> str:
        """'done' | 'do' | 'read' | 'wait' | 'ready' (a wait that has ended)."""
        now = time.time() if now is None else now
        cur = self.current(tube)
        if cur is None:
            return "done"
        if "wait" in cur:
            return "ready" if tube.due is not None and now >= tube.due else "wait"
        return "read" if "read" in cur else "do"

    def _advance(self, tube: Tube, now: float):
        tube.done_at[tube.step] = now
        tube.step += 1
        tube.due = None
        cur = self.current(tube)
        if cur is not None and "wait" in cur:
            frm = cur.get("from", tube.step - 1)
            anchor = tube.done_at.get(frm, now)
            tube.due = anchor + float(cur["wait"])
            if now >= tube.due:            # the wait is already over (e.g. A0 read late)
                pass

    def confirm(self, tube: Tube, now: float | None = None):
        """The operator did the 'do' step, or the wait is over."""
        now = time.time() if now is None else now
        s = self.state(tube, now)
        if s == "wait":
            raise ValueError("Таймер цієї пробірки ще не закінчився")
        if s not in ("do", "ready"):
            raise ValueError("Цей крок — вимір, а не дія")
        self._advance(tube, now)
        # a finished wait right before a read moves on by itself
        if self.state(tube, now) == "ready":
            self._advance(tube, now)

    def record(self, tube: Tube, values: dict, now: float | None = None):
        """values: {wl: A} for the current read step."""
        now = time.time() if now is None else now
        cur = self.current(tube)
        if cur is None or "read" not in cur:
            raise ValueError("Зараз у цієї пробірки не вимір")
        name = cur["read"]
        for wl, a in values.items():
            tube.reads[f"{name}_{int(wl)}"] = float(a)
        if len(cur["wl"]) == 1:
            tube.reads[name] = float(values[cur["wl"][0]])
        self._advance(tube, now)
        if self.state(tube, now) == "ready":
            self._advance(tube, now)

    def tick(self, now: float | None = None) -> list[Tube]:
        """Waits that ended since the last call (to notify); moves them on."""
        now = time.time() if now is None else now
        ended = []
        for t in self.tubes:
            if self.state(t, now) == "ready":
                cur = self.current(t)
                ended.append(t)
                if not cur.get("then"):
                    self._advance(t, now)
        return ended

    def next_tube(self) -> Tube | None:
        """The tube the operator should handle next (ready reads first, then actions)."""
        now = time.time()
        order = {"ready": 0, "read": 1, "do": 2}
        cands = [(order[s], i, t) for i, t in enumerate(self.tubes)
                 if (s := self.state(t, now)) in order]
        zero = self.zero_tube()
        if zero is not None and self.state(zero, now) != "done":
            return zero
        return min(cands, key=lambda c: (c[0], c[1]))[2] if cands else None

    # ----- calculations
    def _value(self, tube: Tube, calc: dict):
        """The 'main' value of a tube: its first computed variable, else its only read."""
        tc = self.p.get("tube_calc", {}).get(tube.type, {})
        if tc:
            return calc.get(next(iter(tc)), NAN)
        for st in self.steps(tube):
            if "read" in st:
                return tube.reads.get(st["read"], NAN) if len(st["wl"]) == 1 else tube.reads.get(
                    f"{st['read']}_{st['wl'][0]}", NAN)
        return NAN

    def _raw(self, tube: Tube):
        """The measured value of a tube: its first read (first wavelength)."""
        for st in self.steps(tube):
            if "read" in st:
                return tube.reads.get(st["read"], tube.reads.get(f"{st['read']}_{st['wl'][0]}", NAN))
        return NAN

    def _lists(self, tubes, calcs) -> dict:
        """`type` = list of measured values, `type.var` = list of a read or computed variable."""
        names = {}
        for t in tubes:
            names.setdefault(t.type, []).append(self._raw(t))
            for k, v in calcs[t.uid].items():
                names.setdefault(f"{t.type}.{k}", []).append(v)
            for k, v in t.reads.items():
                names.setdefault(f"{t.type}.{k}", []).append(v)
        return names

    def compute(self) -> dict:
        """All numbers: series variables, every tube's variables, every sample's results and flags."""
        const = {k: float(v) for k, v in self.p.get("constants", {}).items()}
        tcalc_def = self.p.get("tube_calc", {})
        calcs: dict[str, dict] = {}

        def tube_vars(t: Tube, base: dict):
            names = dict(base)
            names.update(t.reads)
            out = {}
            for var, f in tcalc_def.get(t.type, {}).items():
                v = evaluate(f, names)
                out[var] = v
                names[var] = v
            return out

        series = [t for t in self.tubes if t.sample is None]
        # pass 1: series tubes with constants only, then series variables
        for t in series:
            calcs[t.uid] = tube_vars(t, const)
        svars = dict(const)
        cal = self.p.get("calibration")
        if cal:
            from .calib import fit_calibration
            pts = [(t.x, self._value(t, calcs[t.uid])) for t in series if t.type == cal["type"]]
            pts = [(x, y) for x, y in pts if isnum(y)]
            pre = cal.get("prefix", "cal")
            try:
                f = fit_calibration([p[0] for p in pts], [p[1] for p in pts], cal.get("model", "linear"))
                ys = [p[1] for p in pts]
                svars.update({f"{pre}_slope": f.coef[1], f"{pre}_intercept": f.coef[0], f"{pre}_quad": f.coef[2],
                              f"{pre}_r2": f.r2, f"{pre}_n": float(f.n), f"{pre}_amax": max(ys), f"{pre}_amin": min(ys),
                              f"{pre}_cmax": f.c_max})
                # x from A by the fitted curve (any model): k_conc(A)
                svars[f"{pre}_conc"] = lambda A, _f=f: _f.conc(A)[0] if isnum(A) else NAN
                self.cal_fit = f
            except ValueError:
                svars.update({f"{pre}_slope": NAN, f"{pre}_intercept": NAN, f"{pre}_r2": NAN, f"{pre}_n": float(len(pts)),
                              f"{pre}_conc": lambda A: NAN, f"{pre}_amax": NAN, f"{pre}_amin": NAN})
                self.cal_fit = None
        lists = self._lists(series, calcs)
        for var, f in self.p.get("series_calc", {}).items():
            svars[var] = evaluate(f, {**svars, **lists})
        # series tubes again, now that series variables are known (e.g. GSH control uses Areag)
        for t in series:
            calcs[t.uid] = tube_vars(t, svars)
        lists = self._lists(series, calcs)
        for var, f in self.p.get("series_calc", {}).items():
            svars[var] = evaluate(f, {**svars, **lists})

        samples = []
        for s in self.samples:
            st = [t for t in self.tubes if t.sample == s["code"]]
            base = {**svars, **s["inputs"]}
            # raw values of the sample's tubes are visible to tube formulas (e.g. mean(bg))
            raw = {}
            for t in st:
                raw.setdefault(t.type, []).append(self._raw(t))
            for t in st:
                calcs[t.uid] = tube_vars(t, {**base, **raw})
            names = {**base, **self._lists(st, calcs)}
            res = {}
            for var, f in self.p.get("sample_calc", {}).items():
                v = evaluate(f, names)
                res[var] = v
                names[var] = v
            measured = any(t.reads for t in st if not self.tube_def(t.type).get("zero"))
            flags = self._qc("sample", names) if measured else []
            samples.append({"code": s["code"], "group": s["group"], "inputs": s["inputs"], "res": res,
                            "flags": flags, "measured": measured,
                            "complete": all(self.state(t) == "done" for t in st)})
        series_flags = self._qc("series", {**svars, **lists})
        return {"series": svars, "tubes": calcs, "samples": samples, "series_flags": series_flags}

    def _qc(self, scope: str, names: dict) -> list[str]:
        out = []
        for q in self.p.get("qc", []):
            if q.get("scope", "sample") != scope:
                continue
            v = evaluate(q["expr"], names)
            if not isnum(v):
                continue
            lo, hi = q.get("min"), q.get("max")
            if (lo is not None and v < lo) or (hi is not None and v > hi):
                out.append(q["msg"].format(v=v, pct=v * 100))
        return out

    def group_stats(self, computed: dict) -> list[dict]:
        """Per result: n, mean, SD per group; the change and Welch t-test vs the reference group."""
        groups = []
        for s in computed["samples"]:
            if s["group"] and s["group"] not in groups:
                groups.append(s["group"])
        ref = self.p.get("reference_group") or ("К" if "К" in groups else (groups[0] if groups else None))
        rows = []
        for r in self.p.get("results", []):
            row = {"label": r["label"], "unit": r.get("unit", ""), "groups": {}}
            for g in groups:
                vals = [s["res"].get(r["var"]) for s in computed["samples"] if s["group"] == g]
                vals = _clean([v for v in vals if isnum(v)])
                row["groups"][g] = {"n": len(vals), "mean": _mean(vals), "sd": _sd(vals), "vals": vals}
            for g in groups:
                if g == ref or ref is None:
                    continue
                a, b = row["groups"][ref], row["groups"][g]
                t, df, p = welch(a["vals"], b["vals"])
                from .summary import mann_whitney
                b["p_mw"] = mann_whitney(a["vals"], b["vals"])[1]
                b["change"] = (b["mean"] - a["mean"]) / a["mean"] if a["mean"] else NAN
                b["t"], b["df"], b["p"] = t, df, p
            row["ref"] = ref
            rows.append(row)
        return rows

    # ----- persistence
    def to_json(self) -> dict:
        return {"format": "specrec-protocol-run/1", "protocol": self.p, "started": self.started.isoformat(),
                "samples": self.samples, "tubes": [t.to_dict() for t in self.tubes], "notes": self.notes}

    @classmethod
    def from_json(cls, d: dict) -> "Run":
        r = cls.__new__(cls)
        r.p = d["protocol"]
        r.started = dt.datetime.fromisoformat(d["started"])
        r.samples = d["samples"]
        r.tubes = [Tube(**{k: v for k, v in t.items()}) for t in d["tubes"]]
        for t in r.tubes:
            t.done_at = {int(k): v for k, v in t.done_at.items()}
        r.refs_ready = False
        r.notes = d.get("notes", "")
        return r

    def save(self, path: Path):
        Path(path).write_text(json.dumps(self.to_json(), ensure_ascii=False, indent=1), encoding="utf-8")


def _fmt_x(x) -> str:
    return f"{x:g}".replace(".", ",")


def validate(p: dict) -> list[str]:
    """Problems in a protocol definition (empty list = fine)."""
    errs = []
    for key in ("id", "name", "tubes"):
        if key not in p:
            errs.append(f"немає поля «{key}»")
    if errs:
        return errs
    zeros = [t for t, d in p["tubes"].items() if d.get("zero")]
    if len(zeros) != 1:
        errs.append("потрібна рівно одна пробірка з \"zero\": true (нуль приладу)")
    for t, d in p["tubes"].items():
        for i, st in enumerate(d.get("steps", [])):
            if not any(k in st for k in ("do", "wait", "read")):
                errs.append(f"{t}, крок {i + 1}: має бути do, wait або read")
            if "read" in st and not st.get("wl"):
                errs.append(f"{t}, крок {i + 1}: у read немає довжин хвиль wl")
    for section in ("series_calc", "sample_calc"):
        for var, f in p.get(section, {}).items():
            try:
                ast.parse(str(f), mode="eval")
            except SyntaxError:
                errs.append(f"{section}.{var}: помилка у формулі")
    for t, calc in p.get("tube_calc", {}).items():
        for var, f in calc.items():
            try:
                ast.parse(str(f), mode="eval")
            except SyntaxError:
                errs.append(f"tube_calc.{t}.{var}: помилка у формулі")
    return errs


# ---------------------------------------------------------------- library and import
def protocols_dir() -> Path:
    from . import paths
    d = paths.data_root() / "protocols"
    d.mkdir(parents=True, exist_ok=True)
    return d


def load_library() -> list[tuple[dict, Path | None, list[str]]]:
    """Built-in protocols, then the user's JSON files: (protocol, file or None, problems)."""
    from .protocols_builtin import BUILTIN
    out = [(copy.deepcopy(p), None, []) for p in BUILTIN]
    for f in sorted(protocols_dir().glob("*.json")):
        try:
            p = json.loads(f.read_text(encoding="utf-8"))
            out.append((p, f, validate(p)))
        except Exception as e:      # a broken file must not break the program
            out.append(({"id": f.stem, "name": f"{f.stem} (не читається)", "tubes": {}}, f, [str(e)]))
    return out


def copy_to_user(p: dict) -> Path:
    """Save a copy of a protocol as an editable JSON file; returns its path."""
    q = copy.deepcopy(p)
    base = f"{p['id']}_my"
    i = 1
    while (protocols_dir() / f"{base}{'' if i == 1 else i}.json").exists():
        i += 1
    q["id"] = f"{base}{'' if i == 1 else i}"
    q["name"] = f"{p['name']} — моя копія{'' if i == 1 else ' ' + str(i)}"
    path = protocols_dir() / f"{q['id']}.json"
    path.write_text(json.dumps(q, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def import_samples(path: str | Path) -> list[dict]:
    """Animals from an Excel journal: the sheet «Зразки» (or the first sheet with a «Код» column):
    the header row with «Код» and «Група», then one animal per row until the first empty code."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    sheets = ([wb["Зразки"]] if "Зразки" in wb.sheetnames else []) + list(wb.worksheets)
    for ws in sheets:
        rows = list(ws.iter_rows(values_only=True, max_row=300))
        for i, row in enumerate(rows):
            heads = [str(c).strip().lower() if c is not None else "" for c in row]
            code_col = next((j for j, h in enumerate(heads) if h.startswith("код")), None)
            if code_col is None:
                continue
            grp_col = next((j for j, h in enumerate(heads) if h.startswith("груп")), None)
            out = []
            for r in rows[i + 1:]:
                code = r[code_col] if code_col < len(r) else None
                if code is None or str(code).strip() == "":
                    if out:
                        break
                    continue
                code = str(code).strip()
                if code.lower().startswith("разом"):
                    break
                group = str(r[grp_col]).strip() if grp_col is not None and grp_col < len(r) and r[grp_col] else ""
                out.append({"code": code, "group": group})
            if out:
                return out
    raise ValueError("У файлі не знайдено таблиці з колонкою «Код» (і «Група»)")
