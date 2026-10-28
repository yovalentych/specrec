"""Write the raw readings of protocol series into the lab's Excel journal (a copy).

Only the input cells are filled; the journal's own formulas, flags and the summary sheet
then work as before. Every protocol says where its numbers go ("journal" in the protocol):

    "journal": {
      "sheet": "СОД",                 # sheet of this assay
      "first_row": 17,                # row of the 1st animal; animal i (0-based in «Зразки») → first_row + i
      "sample_cells": {"E": "test#1.A0", "D": "@dilution"},
      "series_cells": {"B8": "control#1.A0"},
      "samples_sheet_cells": {"J": "test#1.fon_380"}   # into «Зразки», row 5 + i
    }

A value is "<tube type>#<n>.<read or computed variable>" or "@<sample input>".
Animals are matched by code with the sheet «Зразки» (codes in column B from row 5); an animal
that is not there is added to the first empty row.
"""
from __future__ import annotations

import math
import re
from pathlib import Path

from .protocol import Run, isnum

SAMPLES_SHEET, SAMPLES_FIRST, SAMPLES_LAST, CODE_COL, GROUP_COL = "Зразки", 5, 60, "B", "C"


def _resolve(run: Run, computed: dict, spec: str, tubes, sample: dict | None):
    if spec.startswith("@"):
        return sample["inputs"].get(spec[1:]) if sample else None
    m = re.match(r"(\w+)#(\d+)\.(\w+)$", spec)
    if not m:
        raise ValueError(f"Незрозуміле поле журналу: {spec}")
    ttype, n, var = m.group(1), int(m.group(2)), m.group(3)
    ts = [t for t in tubes if t.type == ttype]
    if n > len(ts):
        return None
    t = ts[n - 1]
    v = t.reads.get(var, computed["tubes"].get(t.uid, {}).get(var))
    return round(float(v), 4) if isnum(v) else None


def journal_rows(ws) -> dict[str, int]:
    """Animal code -> row in «Зразки»."""
    out = {}
    for r in range(SAMPLES_FIRST, SAMPLES_LAST + 1):
        v = ws[f"{CODE_COL}{r}"].value
        if v is not None and str(v).strip() and not str(v).startswith("="):
            out[str(v).strip()] = r
    return out


def write_runs(journal: str | Path, out: str | Path, runs: list[Run]) -> list[str]:
    """Fill the journal with the runs; save as `out`. Returns a report (lines)."""
    from openpyxl import load_workbook

    wb = load_workbook(journal)          # formulas kept
    if SAMPLES_SHEET not in wb.sheetnames:
        raise ValueError("У журналі немає аркуша «Зразки»")
    ws_s = wb[SAMPLES_SHEET]
    rows = journal_rows(ws_s)
    report = []
    for run in runs:
        spec = run.p.get("journal")
        if not spec:
            report.append(f"{run.p['name']}: у протоколі немає розмітки журналу — пропущено")
            continue
        if spec["sheet"] not in wb.sheetnames:
            report.append(f"{run.p['name']}: у журналі немає аркуша «{spec['sheet']}» — пропущено")
            continue
        ws = wb[spec["sheet"]]
        c = run.compute()
        series = [t for t in run.tubes if t.sample is None]
        n_cells = 0
        for cell, what in spec.get("series_cells", {}).items():
            v = _resolve(run, c, what, series, None)
            if v is not None:
                ws[cell] = v
                n_cells += 1
        added, done = [], 0
        for s in run.samples:
            tubes = [t for t in run.tubes if t.sample == s["code"]]
            if not any(t.reads for t in tubes):
                continue
            r = rows.get(s["code"])
            if r is None:                      # a new animal: first free row of «Зразки»
                r = next((x for x in range(SAMPLES_FIRST, SAMPLES_LAST + 1)
                          if not ws_s[f"{CODE_COL}{x}"].value), None)
                if r is None:
                    report.append(f"  {s['code']}: немає вільного рядка в «Зразки» — пропущено")
                    continue
                ws_s[f"{CODE_COL}{r}"] = s["code"]
                if s.get("group"):
                    ws_s[f"{GROUP_COL}{r}"] = s["group"]
                rows[s["code"]] = r
                added.append(s["code"])
            i = r - SAMPLES_FIRST
            row = spec["first_row"] + i
            for col, what in spec.get("sample_cells", {}).items():
                v = _resolve(run, c, what, tubes, s)
                if v is not None:
                    ws[f"{col}{row}"] = v
                    n_cells += 1
            for col, what in spec.get("samples_sheet_cells", {}).items():
                v = _resolve(run, c, what, tubes, s)
                if v is not None:
                    ws_s[f"{col}{r}"] = v
                    n_cells += 1
            done += 1
        report.append(f"{run.p['name']} → аркуш «{spec['sheet']}»: {done} тварин, {n_cells} клітинок"
                      + (f"; додано в «Зразки»: {', '.join(added)}" if added else ""))
    wb.save(out)
    return report
