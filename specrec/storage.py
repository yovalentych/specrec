"""Saving measurements: every result is written to CSV at once (nothing is lost if the
program closes), plus export of a table to Excel (.xlsx) or CSV on request.

Layout:  SpecRec_data/<YYYY-MM-DD>/photometry.csv
                                  scan_<HHMMSS>_<name>.csv
                                  kinetics_<HHMMSS>_<name>.csv
CSV: comma separator, decimal point, UTF-8 with BOM (opens in Excel, R, Python).
"""
from __future__ import annotations

import csv
import datetime as dt
import math
import re
from pathlib import Path

from . import paths


def day_dir() -> Path:
    d = paths.data_root() / dt.date.today().isoformat()
    d.mkdir(parents=True, exist_ok=True)
    return d


def safe_name(s: str) -> str:
    s = re.sub(r"[^\w\-]+", "_", s.strip(), flags=re.UNICODE).strip("_")
    return s[:40] or "sample"


def fmt(v, digits=4):
    if v is None:
        return ""
    if isinstance(v, float):
        if math.isinf(v):
            return "inf"
        if math.isnan(v):
            return ""
        return f"{v:.{digits}f}"
    return str(v)


def append_row(path: Path, header: list[str], row: list):
    """Append a row; if the file has a different header (older program version), continue
    in a new file next to it (name_2.csv, …) so that the columns never get mixed up."""
    path = Path(path)
    k = 2
    while path.exists():
        with path.open(encoding="utf-8-sig") as f:
            first = next(csv.reader(f), None)
        if first == header:
            break
        stem = path.stem.rsplit("_", 1)[0] if path.stem.rsplit("_", 1)[-1].isdigit() else path.stem
        path = path.with_name(f"{stem}_{k}{path.suffix}")
        k += 1
    new = not path.exists()
    with path.open("a", newline="", encoding="utf-8-sig" if new else "utf-8") as f:
        w = csv.writer(f)
        if new:
            w.writerow(header)
        w.writerow([fmt(v) for v in row])


def write_table(path: Path, header: list[str], rows: list[list], meta: dict | None = None):
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        return write_xlsx(path, header, rows, meta)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        for k, v in (meta or {}).items():
            w.writerow([f"# {k}", v])
        w.writerow(header)
        for r in rows:
            w.writerow([fmt(v) for v in r])
    return path


def write_xlsx(path: Path, header: list[str], rows: list[list], meta: dict | None = None):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Дані"
    r0 = 1
    for k, v in (meta or {}).items():
        ws.cell(r0, 1, k).font = Font(italic=True)
        ws.cell(r0, 2, v)
        r0 += 1
    if meta:
        r0 += 1
    for j, h in enumerate(header, 1):
        ws.cell(r0, j, h).font = Font(bold=True)
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            if isinstance(v, float) and (math.isinf(v) or math.isnan(v)):
                v = "inf" if math.isinf(v) else None
            ws.cell(i, j, v)
    for j, h in enumerate(header, 1):
        ws.column_dimensions[ws.cell(r0, j).column_letter].width = max(10, len(str(h)) + 2)
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    wb.save(path)
    return path


def write_workbook(path: Path, sheets: list[tuple[str, list[str], list[list]]], meta: dict | None = None):
    """Several sheets in one .xlsx: [(title, header, rows), ...]; meta goes on top of the first."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for k, (title, header, rows) in enumerate(sheets):
        ws = wb.create_sheet(title[:31])
        r0 = 1
        if k == 0 and meta:
            for key, v in meta.items():
                ws.cell(r0, 1, key).font = Font(italic=True)
                ws.cell(r0, 2, v)
                r0 += 1
            r0 += 1
        for j, h in enumerate(header, 1):
            ws.cell(r0, j, h).font = Font(bold=True)
            ws.column_dimensions[ws.cell(r0, j).column_letter].width = max(10, min(40, len(str(h)) + 2))
        for i, row in enumerate(rows, r0 + 1):
            for j, v in enumerate(row, 1):
                if isinstance(v, float) and (math.isinf(v) or math.isnan(v)):
                    v = None
                ws.cell(i, j, v)
        ws.freeze_panes = ws.cell(r0 + 1, 1)
    wb.save(path)
    return path
